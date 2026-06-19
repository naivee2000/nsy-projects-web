#!/usr/bin/env python3
"""Build the first-pass DEM claim DISCHARGE section from a PNG/JPG/PDF source."""

from __future__ import annotations

import argparse
import re
import tempfile
from copy import copy
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter, range_boundaries
    from PIL import Image, ImageEnhance, ImageOps
except ModuleNotFoundError as exc:
    missing = exc.name or "a required package"
    raise SystemExit(
        f"Missing Python package: {missing}\n"
        "Install the workbook/OCR helpers with:\n"
        "  python3 -m pip install -r requirements.txt"
    ) from exc

from dem_claim_load_builder import (
    CLEAN_TEMPLATE,
    ORIGINAL_WORKBOOK,
    clear_template_values,
    create_clean_template,
    loose_event_text,
    parse_date,
    parse_time,
    require_tesseract,
    run,
    save_auto_recalc,
    source_to_text_variants,
)

DISCHARGE_BLOCK_START = 39
DISCHARGE_BLOCK_HEIGHT = 25
LOAD_BLOCK_START = 16
LOAD_BLOCK_HEIGHT = 23


DISCHARGE_EVENTS = {
    "nor_tendered": {
        "label": "NOR TENDERED",
        "cell": "D42",
        "patterns": [r"\bnor\b.*tender", r"\binor\b.*tender", r"\b[odi]r\b.*tender"],
        "row_box": (0.284, 0.310),
    },
    "anchor_up": {
        "label": "ANCHOR UP",
        "cell": "D43",
        "patterns": [r"anchor.*aweigh", r"anchor.*aweight", r"anchor.*awe"],
        "row_box": (0.336, 0.362),
    },
    "all_fast": {
        "label": "ALL FAST",
        "cell": "D44",
        "patterns": [r"all\s*made\s*fast", r"allmade\s*fast", r"almade", r"almad", r"\bmade\s*fast"],
        "row_box": (0.435, 0.462),
    },
    "hose_on": {
        "label": "HOSE ON",
        "cell": "D45",
        "patterns": [r"loading\s*arm\s*connect", r"\b[il]?oading\s*arm\s*connect"],
        "range_policy": "end",
        "row_box": (0.690, 0.716),
    },
    "disch_commence": {
        "label": "DISCH COMMENCE",
        "cell": "D46",
        "patterns": [r"commenced\s*unloading", r"\b\w*mmenced\s*unload"],
        "row_box": (0.740, 0.767),
    },
    "disch_complete": {
        "label": "DISCH COMPLETE",
        "cell": "D47",
        "patterns": [r"completed\s+\w*unloading", r"\b\w*pleted\s+\w*unloading"],
        "row_box": (0.764, 0.798),
    },
    "hose_off": {
        "label": "HOSE OFF",
        "cell": "D48",
        "patterns": [r"loading\s*arm\s*disconnect", r"\b[il]?oading\s*arm\s*disconnect"],
        "row_box": (0.889, 0.922),
    },
}


@dataclass(frozen=True)
class Candidate:
    event_key: str
    when: datetime
    variant: str
    line: str


def match_event(line: str) -> str | None:
    normalized = loose_event_text(line)
    for event_key, spec in DISCHARGE_EVENTS.items():
        for pattern in spec["patterns"]:
            if re.search(pattern, normalized):
                return event_key
    return None


def parse_text_candidates(text_variants: list[tuple[str, str]]) -> list[Candidate]:
    candidates: list[Candidate] = []
    for variant, text in text_variants:
        for raw_line in text.splitlines():
            event_key = match_event(raw_line)
            if not event_key:
                continue
            date_parts = parse_date(raw_line)
            if not date_parts:
                continue
            range_policy = DISCHARGE_EVENTS[event_key].get("range_policy", "start")
            hhmm = parse_time(raw_line, range_policy=range_policy)
            if not hhmm:
                continue
            year, month, day = date_parts
            when = datetime(year, month, day, int(hhmm[:2]), int(hhmm[2:]))
            candidates.append(Candidate(event_key, when, variant, raw_line.strip()))
    return candidates


def ocr_crop(image, box: tuple[int, int, int, int], work_dir: Path, stem: str) -> list[str]:
    require_tesseract()
    x1, y1, x2, y2 = box
    crop = image.crop((x1, y1, x2, y2))
    crop = crop.resize((crop.width * 5, crop.height * 5))

    texts: list[str] = []
    variants = []

    gray = ImageOps.autocontrast(crop)
    variants.append(("gray", gray))

    sharp = ImageEnhance.Contrast(gray).enhance(2.4)
    variants.append(("sharp", sharp))

    threshold = sharp.point(lambda p: 255 if p > 175 else 0)
    variants.append(("threshold", threshold))

    for variant_name, variant_image in variants:
        path = work_dir / f"{stem}_{variant_name}.png"
        variant_image.save(path)
        for psm in ("6", "7", "8"):
            result = run(["tesseract", str(path), "stdout", "--psm", psm, "-l", "eng"])
            text = result.stdout.strip()
            if text:
                texts.append(text)
    return texts


def horizontal_line_candidates(image) -> list[int]:
    width, height = image.size
    x1 = round(width * 0.07)
    x2 = round(width * 0.94)
    pixels = image.load()
    centers: list[int] = []

    for dark_threshold in (90, 110, 130, 150, 170, 190, 210):
        counts = [
            sum(1 for x in range(x1, x2) if pixels[x, y] < dark_threshold)
            for y in range(height)
        ]
        for fraction in (0.12, 0.18, 0.24, 0.30, 0.36):
            line_threshold = (x2 - x1) * fraction
            in_cluster = False
            start = 0
            for y, count in enumerate(counts):
                if count > line_threshold and not in_cluster:
                    start = y
                    in_cluster = True
                if in_cluster and (count <= line_threshold or y == height - 1):
                    end = y - 1
                    in_cluster = False
                    center = (start + end) // 2
                    if height * 0.18 < center < height * 0.95:
                        centers.append(center)

    if not centers:
        return []

    centers.sort()
    clustered: list[int] = []
    cluster: list[int] = []
    for center in centers:
        if not cluster or center - cluster[-1] <= 3:
            cluster.append(center)
        else:
            clustered.append(round(sum(cluster) / len(cluster)))
            cluster = [center]
    if cluster:
        clustered.append(round(sum(cluster) / len(cluster)))
    return clustered


def ocr_single_row(image, y1: int, y2: int, work_dir: Path, stem: str) -> list[str]:
    width, _height = image.size
    x1 = round(width * 0.07)
    x2 = round(width * 0.94)
    return ocr_crop(image, (x1, max(0, y1), x2, min(image.height, y2)), work_dir, stem)


def ocr_grid_row(image, y1: int, y2: int, work_dir: Path, stem: str) -> list[str]:
    require_tesseract()
    width, _height = image.size
    x1 = round(width * 0.07)
    x2 = round(width * 0.94)
    crop = image.crop((x1, max(0, y1), x2, min(image.height, y2)))
    crop = crop.resize((crop.width * 4, crop.height * 4))

    variants = []
    gray = ImageOps.autocontrast(crop)
    variants.append(("gray", gray))
    threshold = ImageEnhance.Contrast(gray).enhance(2.0).point(lambda p: 255 if p > 175 else 0)
    variants.append(("threshold", threshold))

    texts: list[str] = []
    for variant_name, variant_image in variants:
        path = work_dir / f"{stem}_{variant_name}.png"
        variant_image.save(path)
        result = run(["tesseract", str(path), "stdout", "--psm", "7", "-l", "eng"])
        text = result.stdout.strip()
        if text:
            texts.append(text)
    return texts


def find_data_start(image, line_candidates: list[int], work_dir: Path) -> tuple[int, float] | None:
    best: tuple[int, float, int, int] | None = None
    _width, height = image.size
    for idx in range(1, len(line_candidates)):
        data_start = line_candidates[idx]
        if not height * 0.23 <= data_start <= height * 0.31:
            continue
        header_found = False
        for prev_idx in range(max(0, idx - 5), idx):
            prev_line = line_candidates[prev_idx]
            if not 14 <= data_start - prev_line <= 38:
                continue
            row_texts = ocr_grid_row(image, prev_line, data_start, work_dir, f"header_{prev_idx}_{idx}")
            header_text = loose_event_text(" ".join(row_texts))
            if "operation" in header_text or ("date" in header_text and "time" in header_text):
                header_found = True
                break
        if not header_found:
            continue

        for step_tenths in range(180, 285, 5):
            step = step_tenths / 10
            matches = 0
            residual = 0
            for row_idx in range(0, 27):
                expected = data_start + row_idx * step
                nearest = min(line_candidates, key=lambda line: abs(line - expected))
                distance = abs(nearest - expected)
                if distance <= max(4.0, step * 0.24):
                    matches += 1
                    residual += round(distance)
            score = matches * 100 - residual
            if best is None or score > best[2] or (score >= best[2] - 150 and data_start < best[0]):
                best = (data_start, step, score, idx)

    if not best:
        return None
    return best[0], best[1]


def image_grid_candidates(source: Path, work_dir: Path) -> list[Candidate]:
    if source.suffix.lower() not in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return []

    image = Image.open(source).convert("L")
    line_candidates = horizontal_line_candidates(image)
    grid = find_data_start(image, line_candidates, work_dir)
    if not grid:
        return []

    data_start, row_height = grid
    candidates: list[Candidate] = []
    width, _height = image.size
    date_time_x1 = round(width * 0.44)
    date_time_x2 = round(width * 0.77)
    for row_idx in range(0, 27):
        y1 = round(data_start + row_idx * row_height) - 2
        y2 = round(data_start + (row_idx + 1) * row_height) + 2
        texts = ocr_grid_row(image, y1, y2, work_dir, f"grid_row_{row_idx}")
        matched_event_keys: set[str] = set()
        for text in texts:
            event_key = match_event(text)
            if not event_key:
                continue
            matched_event_keys.add(event_key)
            date_parts = parse_date(text)
            if not date_parts:
                continue
            range_policy = DISCHARGE_EVENTS[event_key].get("range_policy", "start")
            hhmm = parse_time(text, range_policy=range_policy)
            if not hhmm:
                continue
            year, month, day = date_parts
            when = datetime(year, month, day, int(hhmm[:2]), int(hhmm[2:]))
            candidates.append(Candidate(event_key, when, f"grid_row_{row_idx}", text))

        if row_idx == 1:
            matched_event_keys.add("nor_tendered")
        if row_idx == 3:
            matched_event_keys.add("anchor_up")

        if matched_event_keys:
            cell_texts = ocr_crop(
                image,
                (date_time_x1, y1, date_time_x2, y2),
                work_dir,
                f"grid_cell_{row_idx}",
            )
            for event_key in matched_event_keys:
                for text in cell_texts:
                    date_parts = parse_date(text)
                    if not date_parts:
                        continue
                    range_policy = DISCHARGE_EVENTS[event_key].get("range_policy", "start")
                    hhmm = parse_time(text, range_policy=range_policy)
                    if not hhmm:
                        continue
                    year, month, day = date_parts
                    when = datetime(year, month, day, int(hhmm[:2]), int(hhmm[2:]))
                    candidates.append(Candidate(event_key, when, f"grid_cell_{row_idx}", text))

    return candidates


def image_row_candidates(source: Path, work_dir: Path) -> list[Candidate]:
    if source.suffix.lower() not in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return []

    image = Image.open(source).convert("L")
    width, height = image.size

    date_time_col = (round(width * 0.455), round(width * 0.770))
    candidates: list[Candidate] = []

    for event_key, spec in DISCHARGE_EVENTS.items():
        top_ratio, bottom_ratio = spec["row_box"]
        y1 = max(0, round(height * top_ratio) - 3)
        y2 = min(height, round(height * bottom_ratio) + 3)
        x1, x2 = date_time_col
        texts = ocr_crop(image, (x1, y1, x2, y2), work_dir, f"row_{event_key}")

        for text in texts:
            date_parts = parse_date(text)
            if not date_parts:
                continue
            range_policy = spec.get("range_policy", "start")
            hhmm = parse_time(text, range_policy=range_policy)
            if not hhmm:
                continue
            year, month, day = date_parts
            when = datetime(year, month, day, int(hhmm[:2]), int(hhmm[2:]))
            candidates.append(Candidate(event_key, when, f"row_{event_key}", text))

    return candidates


def choose_event_values(candidates: list[Candidate]) -> tuple[dict[str, datetime], dict[str, list[Candidate]]]:
    grouped: dict[str, list[Candidate]] = defaultdict(list)
    for candidate in candidates:
        grouped[candidate.event_key].append(candidate)

    def preferred_pool(items: list[Candidate]) -> list[Candidate]:
        grid_cell_items = [item for item in items if item.variant.startswith("grid_cell_")]
        grid_items = [item for item in items if item.variant.startswith("grid_row_")]
        text_items = [
            item
            for item in items
            if not item.variant.startswith(("grid_cell_", "grid_row_", "row_"))
        ]
        row_items = [item for item in items if item.variant.startswith("row_")]
        return grid_cell_items or grid_items or text_items or row_items or items

    chosen: dict[str, datetime] = {}
    for event_key, items in grouped.items():
        vote_pool = preferred_pool(items)
        counts = Counter(item.when for item in vote_pool)
        chosen[event_key] = counts.most_common(1)[0][0]

    previous: datetime | None = None
    for event_key in DISCHARGE_EVENTS:
        if event_key not in chosen:
            continue
        if previous and chosen[event_key] < previous:
            vote_pool = preferred_pool(grouped[event_key])
            valid_items = [item for item in vote_pool if item.when >= previous]
            if not valid_items:
                valid_items = [item for item in grouped[event_key] if item.when >= previous]
            if valid_items:
                counts = Counter(item.when for item in valid_items)
                chosen[event_key] = counts.most_common(1)[0][0]
        previous = chosen[event_key]

    return chosen, grouped


def translate_formula(formula: str, origin: str, destination: str) -> str:
    try:
        return Translator(formula, origin=origin).translate_formula(destination)
    except Exception:
        return formula


def copy_row_block(ws, source_start: int, dest_start: int, row_count: int) -> None:
    ws.insert_rows(dest_start, row_count)
    row_offset = dest_start - source_start

    for row in range(source_start, source_start + row_count):
        source_dim = ws.row_dimensions[row]
        dest_dim = ws.row_dimensions[row + row_offset]
        dest_dim.height = source_dim.height
        dest_dim.hidden = source_dim.hidden

        for col in range(1, ws.max_column + 1):
            source = ws.cell(row=row, column=col)
            dest = ws.cell(row=row + row_offset, column=col)
            if isinstance(source.value, str) and source.value.startswith("="):
                dest.value = translate_formula(source.value, source.coordinate, dest.coordinate)
            else:
                dest.value = source.value
            if source.has_style:
                dest._style = copy(source._style)
            if source.number_format:
                dest.number_format = source.number_format
            if source.alignment:
                dest.alignment = copy(source.alignment)
            if source.font:
                dest.font = copy(source.font)
            if source.fill:
                dest.fill = copy(source.fill)
            if source.border:
                dest.border = copy(source.border)

    existing_merges = {str(merged_range) for merged_range in ws.merged_cells.ranges}
    for merged_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if source_start <= min_row and max_row < source_start + row_count:
            new_min_row = min_row + row_offset
            new_max_row = max_row + row_offset
            new_range = (
                f"{get_column_letter(min_col)}{new_min_row}:"
                f"{get_column_letter(max_col)}{new_max_row}"
            )
            if new_range not in existing_merges:
                ws.merge_cells(new_range)
                existing_merges.add(new_range)


def ensure_discharge_blocks(
    ws,
    count: int,
    first_start: int = DISCHARGE_BLOCK_START,
) -> list[int]:
    starts = [first_start]
    for index in range(1, count):
        dest_start = first_start + index * DISCHARGE_BLOCK_HEIGHT
        copy_row_block(ws, first_start, dest_start, DISCHARGE_BLOCK_HEIGHT)
        starts.append(dest_start)
    return starts


def update_total_formulas(ws, discharge_starts: list[int], load_starts: list[int] | None = None) -> None:
    load_starts = load_starts or [LOAD_BLOCK_START]
    allowed_refs = [f"B{start + 19}" for start in load_starts] + [
        f"B{start + 21}" for start in discharge_starts
    ]
    used_refs = [f"D{start + 19}" for start in load_starts] + [
        f"D{start + 21}" for start in discharge_starts
    ]
    dem_refs = [f"B{start + 20}" for start in load_starts] + [
        f"B{start + 22}" for start in discharge_starts
    ]

    def sum_or_blank(refs: list[str]) -> str:
        blank_checks = ",".join(f'{ref}=""' for ref in refs)
        sum_expr = "+".join(f"N({ref})" for ref in refs)
        return f'=IF(AND({blank_checks}),"",{sum_expr})'

    ws["B13"].value = sum_or_blank(allowed_refs)
    ws["C13"].value = sum_or_blank(used_refs)
    ws["D13"].value = '=IF(OR(B13="",C13=""),"",C13-B13)'
    ws["E13"].value = '=IF(OR(D13="",$F$4=""),"",D13/24*$F$4)'

    for start in discharge_starts:
        total_row = start + 24
        for col in range(1, 7):
            ws.cell(total_row, col).value = None

    final_total_row = discharge_starts[-1] + 24
    ws.cell(final_total_row, 1).value = "TTL DEMURRAGE"
    ws.cell(final_total_row, 2).value = sum_or_blank(dem_refs)
    ws.cell(final_total_row, 3).value = "RMB"


def set_discharge_block_formulas(ws, block_start: int) -> None:
    offset = block_start - DISCHARGE_BLOCK_START
    formulas = {
        11: (
            f'=IF(OR(D{42 + offset}="",D{46 + offset}=""),"",'
            f"MIN(D{42 + offset}+TIME(6,0,0),D{46 + offset}))"
        ),
        12: f'=IF(D{48 + offset}="","",D{48 + offset})',
        14: f'=IF(OR(D{51 + offset}="",D{50 + offset}=""),"",(D{51 + offset}-D{50 + offset})*24)',
        21: f'=IF(D{53 + offset}="","",D{53 + offset}-N(D{55 + offset})-N(D{56 + offset})-N(D{57 + offset}))',
        22: f'=IF(OR($F$4="",D{60 + offset}="",B{60 + offset}=""),"",(D{60 + offset}-B{60 + offset})/24*$F$4)',
    }
    ws[f"D{block_start + 11}"].value = formulas[11]
    ws[f"D{block_start + 12}"].value = formulas[12]
    ws[f"D{block_start + 14}"].value = formulas[14]
    ws[f"D{block_start + 21}"].value = formulas[21]
    ws[f"B{block_start + 22}"].value = formulas[22]


def update_discharge_section(ws, events: dict[str, datetime], block_start: int = DISCHARGE_BLOCK_START) -> None:
    row_offset = block_start - DISCHARGE_BLOCK_START
    for event_key, when in events.items():
        cell_ref = DISCHARGE_EVENTS[event_key]["cell"]
        row = int(cell_ref[1:]) + row_offset
        cell_ref = f"D{row}"
        ws[cell_ref].value = when
        ws[cell_ref].number_format = "m/d/yy h:mm"
    set_discharge_block_formulas(ws, block_start)


def save_workbook(
    template: Path,
    output: Path,
    event_sets: list[dict[str, datetime]],
    sheet_name: str | None,
    load_starts: list[int] | None = None,
) -> None:
    workbook = load_workbook(template)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    clear_template_values(worksheet)
    block_starts = ensure_discharge_blocks(worksheet, len(event_sets))
    for block_start, events in zip(block_starts, event_sets):
        update_discharge_section(worksheet, events, block_start)
    update_total_formulas(worksheet, block_starts, load_starts)
    save_auto_recalc(workbook, output)


def print_report(events: dict[str, datetime], grouped: dict[str, list[Candidate]]) -> None:
    print("Extracted DISCHARGE events:")
    for event_key, spec in DISCHARGE_EVENTS.items():
        when = events.get(event_key)
        if not when:
            print(f"  - {spec['label']}: MISSING")
            continue
        votes = Counter(candidate.when for candidate in grouped[event_key])[when]
        print(f"  - {spec['label']}: {when:%Y-%m-%d %H:%M} ({votes} OCR vote(s))")


def extract_discharge_events(source: Path, work_dir: Path) -> tuple[dict[str, datetime], dict[str, list[Candidate]]]:
    text_variants = source_to_text_variants(source, work_dir)
    for name, text in text_variants:
        (work_dir / f"{source.stem}_{name}.txt").write_text(text, encoding="utf-8")

    candidates = (
        parse_text_candidates(text_variants)
        + image_grid_candidates(source, work_dir)
        + image_row_candidates(source, work_dir)
    )
    return choose_event_values(candidates)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Populate the DISCHARGE section of a DEM claim workbook.")
    parser.add_argument("--template", default=CLEAN_TEMPLATE, help="Clean Excel template path")
    parser.add_argument(
        "--template-source",
        default=ORIGINAL_WORKBOOK,
        help="Filled workbook to clean when the clean template needs to be created",
    )
    parser.add_argument(
        "--source",
        nargs="+",
        default=["Discharge1.png"],
        help="One or more discharge SOF source image/PDF files, in workbook order",
    )
    parser.add_argument("--output", default="outputs/discharge1_section_from_png.xlsx", help="Output workbook path")
    parser.add_argument("--sheet", default=None, help="Worksheet name; defaults to the active sheet")
    parser.add_argument("--debug-dir", default=None, help="Keep OCR crops/text in this folder")
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    template = Path(args.template).expanduser().resolve()
    template_source = Path(args.template_source).expanduser().resolve()
    sources = [Path(source).expanduser().resolve() for source in args.source]
    output = Path(args.output).expanduser().resolve()

    if not template.exists():
        if not template_source.exists():
            raise FileNotFoundError(template_source)
        create_clean_template(template_source, template, args.sheet)
        print(f"Created clean template: {template}")

    for source in sources:
        if not source.exists():
            raise FileNotFoundError(source)

    event_sets: list[dict[str, datetime]] = []
    grouped_sets: list[dict[str, list[Candidate]]] = []
    if args.debug_dir:
        work_dir = Path(args.debug_dir).expanduser().resolve()
        work_dir.mkdir(parents=True, exist_ok=True)
        for source in sources:
            source_work_dir = work_dir / source.stem
            source_work_dir.mkdir(parents=True, exist_ok=True)
            events, grouped = extract_discharge_events(source, source_work_dir)
            event_sets.append(events)
            grouped_sets.append(grouped)
    else:
        with tempfile.TemporaryDirectory(prefix="dem_claim_discharge_ocr_") as tmp:
            root_work_dir = Path(tmp)
            for source in sources:
                source_work_dir = root_work_dir / source.stem
                source_work_dir.mkdir(parents=True, exist_ok=True)
                events, grouped = extract_discharge_events(source, source_work_dir)
                event_sets.append(events)
                grouped_sets.append(grouped)

    for source, events, grouped in zip(sources, event_sets, grouped_sets):
        missing = [
            spec["label"]
            for key, spec in DISCHARGE_EVENTS.items()
            if spec.get("required", True) and key not in events
        ]
        if missing:
            print(f"Source: {source.name}")
            print_report(events, grouped)
            raise RuntimeError(f"OCR could not find required DISCHARGE events: {', '.join(missing)}")

    save_workbook(template, output, event_sets, args.sheet)
    for source, events, grouped in zip(sources, event_sets, grouped_sets):
        print(f"Source: {source.name}")
        print_report(events, grouped)
    print(f"Saved workbook: {output}")


if __name__ == "__main__":
    main()
