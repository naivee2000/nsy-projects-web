#!/usr/bin/env python3
"""Build a DEM claim workbook from load and discharge SOF image/PDF sources."""

from __future__ import annotations

import argparse
import tempfile
from collections import Counter
from copy import copy
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    missing = exc.name or "a required package"
    raise SystemExit(
        f"Missing Python package: {missing}\n"
        "Install the workbook/OCR helpers with:\n"
        "  python3 -m pip install -r requirements.txt"
    ) from exc

from dem_claim_discharge_builder import (
    DISCHARGE_BLOCK_HEIGHT,
    DISCHARGE_BLOCK_START,
    DISCHARGE_EVENTS,
    ensure_discharge_blocks,
    copy_row_block,
    extract_discharge_events,
    set_discharge_block_formulas,
    update_discharge_section,
    update_total_formulas,
)
from dem_claim_load_builder import (
    CLEAN_TEMPLATE,
    LOAD_EVENTS,
    ORIGINAL_WORKBOOK,
    clear_template_values,
    create_clean_template,
    extract_load_events,
    save_auto_recalc,
)

LOAD_BLOCK_START = 16
LOAD_BLOCK_HEIGHT = 23
DOC_DATE_BASE_ROW = 65
DOC_DATE_LABEL_COL = 5
DOC_DATE_VALUE_COL = 6
DOC_DATE_FORMAT = "m/d/yy"


def document_date_row(discharge_starts: list[int]) -> int:
    final_total_row = discharge_starts[-1] + DISCHARGE_BLOCK_HEIGHT - 1
    return final_total_row + 2


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.font:
        target.font = copy(source.font)
    if source.fill:
        target.fill = copy(source.fill)
    if source.border:
        target.border = copy(source.border)


def set_document_date(ws, row: int | None = None, value: date | datetime | None = None) -> None:
    target_row = row or DOC_DATE_BASE_ROW
    source_label = ws.cell(DOC_DATE_BASE_ROW, DOC_DATE_LABEL_COL)
    source_value = ws.cell(DOC_DATE_BASE_ROW, DOC_DATE_VALUE_COL)
    target_label = ws.cell(target_row, DOC_DATE_LABEL_COL)
    target_value = ws.cell(target_row, DOC_DATE_VALUE_COL)

    base_still_has_footer = source_label.value == "Doc Date:"
    if target_row != DOC_DATE_BASE_ROW and base_still_has_footer:
        copy_cell_style(source_label, target_label)
        copy_cell_style(source_value, target_value)
        source_label.value = None
        source_value.value = None

    target_label.value = "Doc Date:"
    target_value.value = value or date.today()
    target_value.number_format = DOC_DATE_FORMAT


def ensure_load_blocks(ws, count: int) -> list[int]:
    starts = [LOAD_BLOCK_START]
    for index in range(1, count):
        dest_start = LOAD_BLOCK_START + index * LOAD_BLOCK_HEIGHT
        copy_row_block(ws, LOAD_BLOCK_START, dest_start, LOAD_BLOCK_HEIGHT)
        starts.append(dest_start)
    return starts


def set_load_block_formulas(ws, block_start: int) -> None:
    offset = block_start - LOAD_BLOCK_START
    ws[f"D{27 + offset}"].value = (
        f'=IF(OR(D{19 + offset}="",D{23 + offset}=""),"",'
        f"MIN(D{19 + offset}+TIME(6,0,0),D{23 + offset}))"
    )
    ws[f"D{28 + offset}"].value = f'=IF(D{25 + offset}="","",D{25 + offset})'
    ws[f"D{30 + offset}"].value = (
        f'=IF(OR(D{28 + offset}="",D{27 + offset}=""),"",'
        f"(D{28 + offset}-D{27 + offset})*24)"
    )
    ws[f"D{32 + offset}"].value = (
        f'=IF(OR(D{21 + offset}="",D{20 + offset}=""),"",'
        f"(D{21 + offset}-D{20 + offset})*24)"
    )
    ws[f"D{35 + offset}"].value = (
        f'=IF(D{30 + offset}="","",D{30 + offset}-N(D{32 + offset})+N(D{33 + offset}))'
    )
    ws[f"B{36 + offset}"].value = (
        f'=IF(OR($F$4="",D{35 + offset}="",B{35 + offset}=""),"",'
        f"(D{35 + offset}-B{35 + offset})/24*$F$4)"
    )


def update_load_section(ws, events: dict[str, datetime], block_start: int) -> None:
    row_offset = block_start - LOAD_BLOCK_START
    for event_key, when in events.items():
        cell_ref = LOAD_EVENTS[event_key]["cell"]
        row = int(cell_ref[1:]) + row_offset
        target = f"D{row}"
        ws[target].value = when
        ws[target].number_format = "m/d/yy h:mm"
    set_load_block_formulas(ws, block_start)


def report_events(
    section: str,
    specs: dict,
    events: dict[str, datetime],
    grouped: dict,
) -> None:
    print(f"Extracted {section} events:")
    for event_key, spec in specs.items():
        when = events.get(event_key)
        if not when:
            print(f"  - {spec['label']}: MISSING")
            continue
        votes = Counter(candidate.when for candidate in grouped[event_key])[when]
        print(f"  - {spec['label']}: {when:%Y-%m-%d %H:%M} ({votes} OCR vote(s))")


def required_missing(specs: dict, events: dict[str, datetime]) -> list[str]:
    return [
        spec["label"]
        for event_key, spec in specs.items()
        if spec.get("required", True) and event_key not in events
    ]


def extract_sources(
    sources: list[Path],
    section: str,
    extractor,
    specs: dict,
    work_root: Path,
) -> list[tuple[Path, dict[str, datetime], dict]]:
    extracted = []
    for index, source in enumerate(sources, start=1):
        if not source.exists():
            raise FileNotFoundError(source)
        source_work_dir = work_root / f"{index:02d}_{section}_{source.stem}"
        source_work_dir.mkdir(parents=True, exist_ok=True)
        events, grouped = extractor(source, source_work_dir)
        missing = required_missing(specs, events)
        if missing:
            print(f"Source: {source.name}")
            report_events(section.upper(), specs, events, grouped)
            raise RuntimeError(
                f"OCR could not find required {section.upper()} events: {', '.join(missing)}"
            )
        extracted.append((source, events, grouped))
    return extracted


def save_workbook(
    template: Path,
    output: Path,
    load_event_sets: list[dict[str, datetime]],
    discharge_event_sets: list[dict[str, datetime]],
    sheet_name: str | None,
) -> tuple[list[int], list[int]]:
    workbook = load_workbook(template)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    clear_template_values(worksheet)

    load_count = max(1, len(load_event_sets))
    load_starts = ensure_load_blocks(worksheet, load_count)
    for block_start in load_starts:
        set_load_block_formulas(worksheet, block_start)
    for block_start, events in zip(load_starts, load_event_sets):
        update_load_section(worksheet, events, block_start)

    first_discharge_start = DISCHARGE_BLOCK_START + (load_count - 1) * LOAD_BLOCK_HEIGHT
    discharge_count = max(1, len(discharge_event_sets))
    discharge_starts = ensure_discharge_blocks(
        worksheet,
        discharge_count,
        first_start=first_discharge_start,
    )
    for block_start in discharge_starts:
        set_discharge_block_formulas(worksheet, block_start)
    for block_start, events in zip(discharge_starts, discharge_event_sets):
        update_discharge_section(worksheet, events, block_start)

    update_total_formulas(worksheet, discharge_starts, load_starts)
    set_document_date(worksheet, row=document_date_row(discharge_starts))
    save_auto_recalc(workbook, output)
    return load_starts, discharge_starts


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Populate a DEM claim workbook top-down from SOF sources.")
    parser.add_argument("--template", default=CLEAN_TEMPLATE, help="Clean Excel template path")
    parser.add_argument(
        "--template-source",
        default=ORIGINAL_WORKBOOK,
        help="Filled workbook to clean when the clean template needs to be created",
    )
    parser.add_argument(
        "--load-source",
        nargs="*",
        default=["Load1.png"],
        help="Zero or more load SOF source image/PDF files, in workbook order",
    )
    parser.add_argument(
        "--discharge-source",
        nargs="*",
        default=["Discharge1.png"],
        help="Zero or more discharge SOF source image/PDF files, in workbook order",
    )
    parser.add_argument("--output", default="outputs/dem_claim_from_sources.xlsx", help="Output workbook path")
    parser.add_argument("--sheet", default=None, help="Worksheet name; defaults to the active sheet")
    parser.add_argument("--debug-dir", default=None, help="Keep OCR crops/text in this folder")
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    template = Path(args.template).expanduser().resolve()
    template_source = Path(args.template_source).expanduser().resolve()
    load_sources = [Path(source).expanduser().resolve() for source in args.load_source]
    discharge_sources = [Path(source).expanduser().resolve() for source in args.discharge_source]
    output = Path(args.output).expanduser().resolve()

    if not template.exists():
        if not template_source.exists():
            raise FileNotFoundError(template_source)
        create_clean_template(template_source, template, args.sheet)
        print(f"Created clean template: {template}")

    if args.debug_dir:
        work_root = Path(args.debug_dir).expanduser().resolve()
        work_root.mkdir(parents=True, exist_ok=True)
        load_sets = extract_sources(load_sources, "load", extract_load_events, LOAD_EVENTS, work_root)
        discharge_sets = extract_sources(
            discharge_sources,
            "discharge",
            extract_discharge_events,
            DISCHARGE_EVENTS,
            work_root,
        )
    else:
        with tempfile.TemporaryDirectory(prefix="dem_claim_workflow_ocr_") as tmp:
            work_root = Path(tmp)
            load_sets = extract_sources(load_sources, "load", extract_load_events, LOAD_EVENTS, work_root)
            discharge_sets = extract_sources(
                discharge_sources,
                "discharge",
                extract_discharge_events,
                DISCHARGE_EVENTS,
                work_root,
            )

    load_event_sets = [events for _source, events, _grouped in load_sets]
    discharge_event_sets = [events for _source, events, _grouped in discharge_sets]
    load_starts, discharge_starts = save_workbook(
        template,
        output,
        load_event_sets,
        discharge_event_sets,
        args.sheet,
    )

    for source, events, grouped in load_sets:
        print(f"Source: {source.name}")
        report_events("LOAD", LOAD_EVENTS, events, grouped)
    for source, events, grouped in discharge_sets:
        print(f"Source: {source.name}")
        report_events("DISCHARGE", DISCHARGE_EVENTS, events, grouped)
    print(f"Load block starts: {', '.join(str(row) for row in load_starts)}")
    print(f"Discharge block starts: {', '.join(str(row) for row in discharge_starts)}")
    print(f"Saved workbook: {output}")


if __name__ == "__main__":
    main()
