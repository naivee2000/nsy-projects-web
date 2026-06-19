#!/usr/bin/env python3
"""Build the first-pass DEM claim LOAD section from a PNG/JPG/PDF source.

This script keeps the existing Excel workbook as the layout template. It clears
the editable top DEM CLAIM values, OCRs the load statement-of-facts table, and
updates the LOAD event dates/times in the lower section.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
    from PIL import Image, ImageEnhance, ImageOps
    from pypdf import PdfReader
except ModuleNotFoundError as exc:
    missing = exc.name or "a required package"
    raise SystemExit(
        f"Missing Python package: {missing}\n"
        "Install the workbook/OCR helpers with:\n"
        "  python3 -m pip install -r requirements.txt"
    ) from exc


MONTHS = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "mey": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}

ORIGINAL_WORKBOOK = "MT HONG SUI V.2609(-066J) DEM CLAIM.xlsx"
CLEAN_TEMPLATE = "DEM_CLAIM_TEMPLATE.xlsx"

LOAD_EVENTS = {
    "nor_tendered": {
        "label": "NOR TENDERED",
        "cell": "D19",
        "patterns": [r"\bnor\b.*tender", r"\binor\b.*tender", r"\b[odi]r\b.*tender"],
    },
    "anchor_up": {
        "label": "ANCHOR UP",
        "cell": "D20",
        "patterns": [r"anchor.*aweigh", r"anchor.*aweight"],
        "required": False,
    },
    "all_fast": {
        "label": "ALL FAST",
        "cell": "D21",
        "patterns": [r"all\s*made\s*fast", r"almade\s*fast", r"allmade\s*fast", r"\bmade\s*fast"],
    },
    "hose_on": {
        "label": "HOSE ON",
        "cell": "D22",
        "patterns": [r"loading\s*arm\s*connect", r"\bading\s*arm\s*connect"],
    },
    "load_commence": {
        "label": "LOAD COMMENCE",
        "cell": "D23",
        "patterns": [r"commenced\s*loading", r"commenced.*load", r"\b\w?mmenced\s*loading"],
    },
    "load_complete": {
        "label": "LOAD COMPLETE",
        "cell": "D24",
        "patterns": [r"completed\s+\w*oading", r"completed.*load", r"\b\w*pleted\s+\w*oading"],
    },
    "hose_off": {
        "label": "HOSE OFF",
        "cell": "D25",
        "patterns": [r"loading\s*arm\s*disconnect", r"\b[fl]?ading\s*arm\s*disconnect"],
        "range_policy": "end",
    },
}

TOP_INPUT_CELLS = [
    "B2",
    "B3",
    "B4",
    "B5",
    "B6",
    "F2",
    "F3",
    "F4",
    "F5",
    "F6",
    "D9",
    "D10",
]

LOAD_VALUE_CELLS = [
    "C16",
    "E16",
    "F16",
    "D19",
    "D20",
    "D21",
    "D22",
    "D23",
    "D24",
    "D25",
    "D27",
    "D28",
    "E27",
    "E28",
    "D33",
    "B35",
]

DISCHARGE_VALUE_CELLS = [
    "C39",
    "E39",
    "C40",
    "F40",
    "D42",
    "D43",
    "D44",
    "D45",
    "D46",
    "D47",
    "D48",
    "E42",
    "E43",
    "E44",
    "E45",
    "E46",
    "E47",
    "E48",
    "D50",
    "D51",
    "E50",
    "E51",
    "D55",
    "D56",
    "D57",
    "D58",
    "B60",
]

FORMULAS = {
    "B13": '=IF(AND(B35="",B60=""),"",N(B35)+N(B60))',
    "C13": '=IF(AND(D35="",D60=""),"",N(D35)+N(D60))',
    "D13": '=IF(OR(B13="",C13=""),"",C13-B13)',
    "E13": '=IF(OR(D13="",$F$4=""),"",D13/24*$F$4)',
    "D27": '=IF(OR(D19="",D23=""),"",MIN(D19+TIME(6,0,0),D23))',
    "D28": '=IF(D25="","",D25)',
    "D30": '=IF(OR(D28="",D27=""),"",(D28-D27)*24)',
    "D32": '=IF(OR(D21="",D20=""),"",(D21-D20)*24)',
    "D35": '=IF(D30="","",D30-N(D32)+N(D33))',
    "B36": '=IF(OR($F$4="",D35="",B35=""),"",(D35-B35)/24*$F$4)',
    "D50": '=IF(OR(D42="",D46=""),"",MIN(D42+TIME(6,0,0),D46))',
    "D51": '=IF(D48="","",D48)',
    "D53": '=IF(OR(D51="",D50=""),"",(D51-D50)*24)',
    "D60": '=IF(D53="","",D53-N(D55)-N(D56)-N(D57))',
    "F60": '=IF(OR(C40="",D47="",D46=""),"",C40/((D47-D46)*24-10.5))',
    "B61": '=IF(OR($F$4="",D60="",B60=""),"",(D60-B60)/24*$F$4)',
    "B63": '=IF(AND(B36="",B61=""),"",N(B36)+N(B61))',
}


@dataclass(frozen=True)
class Candidate:
    event_key: str
    when: datetime
    variant: str
    line: str


def run(cmd: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(cmd, text=True, capture_output=True, check=True)


def require_tesseract() -> str:
    path = shutil.which("tesseract")
    if not path:
        raise RuntimeError("tesseract was not found on PATH. Install Tesseract OCR first.")
    return path


def preprocess_image_variants(source: Path, work_dir: Path) -> list[tuple[str, Path, int]]:
    """Create OCR-friendly crops. Coordinates are ratios for this SOF form."""
    image = Image.open(source).convert("L")
    width, height = image.size
    crop_box = (
        round(width * 0.095),
        round(height * 0.23),
        round(width * 0.92),
        round(height * 0.91),
    )
    table = image.crop(crop_box)
    scaled = table.resize((table.width * 3, table.height * 3))

    variants: list[tuple[str, Path, int]] = []

    gray = ImageOps.autocontrast(scaled)
    gray_path = work_dir / "load_table_gray.png"
    gray.save(gray_path)
    variants.append(("gray_psm11", gray_path, 11))

    for threshold in (145, 155, 165, 175, 185, 195, 205):
        img = ImageEnhance.Contrast(scaled).enhance(1.8)
        img = ImageOps.autocontrast(img)
        img = img.point(lambda p, th=threshold: 255 if p > th else 0)
        out = work_dir / f"load_table_threshold_{threshold}.png"
        img.save(out)
        variants.append((f"threshold_{threshold}_psm6", out, 6))

    return variants


def image_to_text_variants(source: Path, work_dir: Path) -> list[tuple[str, str]]:
    require_tesseract()
    outputs: list[tuple[str, str]] = []
    for name, image_path, psm in preprocess_image_variants(source, work_dir):
        result = run(["tesseract", str(image_path), "stdout", "--psm", str(psm), "-l", "eng"])
        outputs.append((name, result.stdout))
    return outputs


def pdf_to_text_or_image_variants(source: Path, work_dir: Path) -> list[tuple[str, str]]:
    reader = PdfReader(str(source))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if "Operations" in text or "NOR" in text or "loading" in text.lower():
        return [("pdf_text", text)]

    sips = shutil.which("sips")
    if not sips:
        raise RuntimeError(
            "This PDF appears image-based and no PDF-to-image converter was found. "
            "On macOS, the script can use sips when it is available."
        )

    png_path = work_dir / "pdf_page_1.png"
    run([sips, "-s", "format", "png", str(source), "--out", str(png_path)])
    return image_to_text_variants(png_path, work_dir)


def source_to_text_variants(source: Path, work_dir: Path) -> list[tuple[str, str]]:
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return pdf_to_text_or_image_variants(source, work_dir)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return image_to_text_variants(source, work_dir)
    raise ValueError(f"Unsupported source type: {source.suffix}")


def loose_event_text(text: str) -> str:
    text = text.lower()
    text = text.replace("|", " ")
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def match_event(line: str) -> str | None:
    normalized = loose_event_text(line)
    for event_key, spec in LOAD_EVENTS.items():
        for pattern in spec["patterns"]:
            if re.search(pattern, normalized):
                return event_key
    return None


def normalize_ocr_dates(text: str) -> str:
    normalized = text
    normalized = normalized.replace("，", ",")
    normalized = re.sub(r"(?i)zoz[e6g]", "2026", normalized)
    normalized = re.sub(r"(?i)20z6", "2026", normalized)
    normalized = re.sub(r"(?i)2o26", "2026", normalized)
    normalized = re.sub(r"(?i)20(?:0|2)8", "2026", normalized)
    normalized = re.sub(r"(?i)20(?:0|2)6", "2026", normalized)
    normalized = re.sub(r"(?i)mey|mav", "May", normalized)
    return normalized


def parse_date(line: str) -> tuple[int, int, int] | None:
    normalized = normalize_ocr_dates(line)
    month_names = "Jan|Feb|Mar|Apr|May|Mey|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
    fuzzy = re.compile(
        rf"(?P<day>[0-9OoIlTt]{{1,4}})\s*[\._,\-\s]*"
        rf"(?P<month>{month_names})\s*[\._,\-\s]*(?P<year>20\d{{2}})",
        re.IGNORECASE,
    )
    for match in fuzzy.finditer(normalized):
        day_text = match.group("day").translate(
            str.maketrans({"O": "0", "o": "0", "I": "1", "l": "1", "i": "1", "T": "1", "t": "1"})
        )
        digits = re.sub(r"\D", "", day_text)
        if not digits:
            continue
        day_options = []
        if len(digits) >= 2:
            last_two = int(digits[-2:])
            if last_two == 40:
                day_options.append(10)
            day_options.append(last_two)
        day_options.append(int(digits[-1:]))
        for day in day_options:
            if 1 <= day <= 31:
                month = MONTHS[match.group("month").lower()[:3]]
                return int(match.group("year")), month, day

    pattern = re.compile(
        r"(?<!\d)(?P<day>[0-3]?\d)\s*[\._,\-\s]*"
        r"(?P<month>Jan|Feb|Mar|Apr|May|Mey|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
        r"\s*[\._,\-\s]*(?P<year>20\d{2})",
        re.IGNORECASE,
    )
    match = pattern.search(normalized)
    if not match:
        return None

    day = int(match.group("day"))
    month = MONTHS[match.group("month").lower()[:3]]
    year = int(match.group("year"))
    if not 1 <= day <= 31:
        return None
    return year, month, day


def digitish_to_digits(token: str) -> str:
    table = str.maketrans(
        {
            "O": "0",
            "o": "0",
            "D": "0",
            "Q": "0",
            "I": "1",
            "l": "1",
            "i": "1",
            "S": "3",
            "s": "3",
            "B": "8",
            "a": "8",
            "g": "9",
            "q": "9",
            "Z": "2",
            "z": "2",
        }
    )
    return token.translate(table)


def valid_hhmm(value: str) -> str | None:
    digits = re.sub(r"\D", "", value)
    if len(digits) == 3:
        digits = "0" + digits
    if len(digits) != 4:
        return None
    hour = int(digits[:2])
    minute = int(digits[2:])
    if 0 <= hour <= 23 and 0 <= minute <= 59:
        return digits
    return None


def parse_time(line: str, range_policy: str = "start") -> str | None:
    text = normalize_ocr_dates(line)
    text = digitish_to_digits(text)

    range_matches: list[tuple[str, str]] = []
    for match in re.finditer(r"(?<!\d)([0-2]?\d{3})\s*[-–—_]\s*([0-2]?\d{3})(?!\d)", text):
        start = valid_hhmm(match.group(1))
        end = valid_hhmm(match.group(2))
        if start and end:
            range_matches.append((start, end))

    for match in re.finditer(r"(?<!\d)([0-2]\d{7})(?!\d)", text):
        start = valid_hhmm(match.group(1)[:4])
        end = valid_hhmm(match.group(1)[4:])
        if start and end:
            range_matches.append((start, end))

    if range_matches:
        start, end = range_matches[-1]
        return end if range_policy == "end" else start

    singles: list[str] = []
    for match in re.finditer(r"(?<!\d)([0-2]?\d{3})(?!\d)", text):
        hhmm = valid_hhmm(match.group(1))
        if hhmm and hhmm not in {"2026"}:
            singles.append(hhmm)
    return singles[-1] if singles else None


def parse_candidates(text_variants: list[tuple[str, str]]) -> list[Candidate]:
    candidates: list[Candidate] = []
    for variant, text in text_variants:
        for raw_line in text.splitlines():
            event_key = match_event(raw_line)
            if not event_key:
                continue
            date_parts = parse_date(raw_line)
            if not date_parts:
                continue
            range_policy = LOAD_EVENTS[event_key].get("range_policy", "start")
            hhmm = parse_time(raw_line, range_policy=range_policy)
            if not hhmm:
                continue
            year, month, day = date_parts
            when = datetime(year, month, day, int(hhmm[:2]), int(hhmm[2:]))
            candidates.append(Candidate(event_key, when, variant, raw_line.strip()))
    return candidates


def choose_event_values(candidates: list[Candidate]) -> tuple[dict[str, datetime], dict[str, list[Candidate]]]:
    grouped: dict[str, list[Candidate]] = defaultdict(list)
    for candidate in candidates:
        grouped[candidate.event_key].append(candidate)

    chosen: dict[str, datetime] = {}
    for event_key, items in grouped.items():
        counts = Counter(item.when for item in items)
        chosen[event_key] = counts.most_common(1)[0][0]
    return chosen, grouped


def set_guarded_formulas(ws) -> None:
    for ref, formula in FORMULAS.items():
        ws[ref].value = formula


def clear_top_dem_claim(ws) -> None:
    for ref in TOP_INPUT_CELLS:
        ws[ref].value = None


def clear_load_section_values(ws) -> None:
    for ref in LOAD_VALUE_CELLS:
        ws[ref].value = None


def clear_discharge_section(ws) -> None:
    for ref in DISCHARGE_VALUE_CELLS:
        ws[ref].value = None


def clear_template_values(ws) -> None:
    clear_top_dem_claim(ws)
    clear_load_section_values(ws)
    clear_discharge_section(ws)
    set_guarded_formulas(ws)


def update_load_section(ws, events: dict[str, datetime]) -> None:
    for event_key, when in events.items():
        cell_ref = LOAD_EVENTS[event_key]["cell"]
        ws[cell_ref].value = when
        ws[cell_ref].number_format = "m/d/yy h:mm"

    set_guarded_formulas(ws)


def save_auto_recalc(wb, output: Path) -> None:
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def create_clean_template(source_template: Path, clean_template: Path, sheet_name: str | None = None) -> None:
    wb = load_workbook(source_template)
    ws = wb[sheet_name] if sheet_name else wb.active
    clear_template_values(ws)
    save_auto_recalc(wb, clean_template)


def save_workbook(template: Path, output: Path, events: dict[str, datetime], args: argparse.Namespace) -> None:
    wb = load_workbook(template)
    ws = wb[args.sheet] if args.sheet else wb.active

    clear_template_values(ws)
    update_load_section(ws, events)
    save_auto_recalc(wb, output)


def print_report(events: dict[str, datetime], grouped: dict[str, list[Candidate]]) -> None:
    print("Extracted LOAD events:")
    for event_key, spec in LOAD_EVENTS.items():
        when = events.get(event_key)
        if not when:
            print(f"  - {spec['label']}: MISSING")
            continue
        votes = Counter(candidate.when for candidate in grouped[event_key])[when]
        print(f"  - {spec['label']}: {when:%Y-%m-%d %H:%M} ({votes} OCR vote(s))")


def extract_load_events(source: Path, work_dir: Path) -> tuple[dict[str, datetime], dict[str, list[Candidate]]]:
    text_variants = source_to_text_variants(source, work_dir)
    for name, text in text_variants:
        (work_dir / f"{source.stem}_{name}.txt").write_text(text, encoding="utf-8")

    candidates = parse_candidates(text_variants)
    return choose_event_values(candidates)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Populate the LOAD section of a DEM claim workbook.")
    parser.add_argument("--template", default=CLEAN_TEMPLATE, help="Clean Excel template path")
    parser.add_argument(
        "--template-source",
        default=ORIGINAL_WORKBOOK,
        help="Filled workbook to clean when the clean template needs to be created",
    )
    parser.add_argument(
        "--create-template",
        action="store_true",
        help="Only create/update the clean template, then exit",
    )
    parser.add_argument("--source", default="Load1.png", help="Load SOF source image/PDF")
    parser.add_argument("--output", default="outputs/load_section_from_png.xlsx", help="Output workbook path")
    parser.add_argument("--sheet", default=None, help="Worksheet name; defaults to the active sheet")
    parser.add_argument("--debug-dir", default=None, help="Keep OCR crops/text in this folder")
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    template = Path(args.template).expanduser().resolve()
    template_source = Path(args.template_source).expanduser().resolve()
    source = Path(args.source).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()

    if args.create_template:
        if not template_source.exists():
            raise FileNotFoundError(template_source)
        create_clean_template(template_source, template, args.sheet)
        print(f"Created clean template: {template}")
        return

    if not template.exists():
        if not template_source.exists():
            raise FileNotFoundError(template_source)
        create_clean_template(template_source, template, args.sheet)
        print(f"Created clean template: {template}")

    if not source.exists():
        raise FileNotFoundError(source)

    if args.debug_dir:
        work_dir = Path(args.debug_dir).expanduser().resolve()
        work_dir.mkdir(parents=True, exist_ok=True)
        text_variants = source_to_text_variants(source, work_dir)
        for name, text in text_variants:
            (work_dir / f"{name}.txt").write_text(text, encoding="utf-8")
    else:
        with tempfile.TemporaryDirectory(prefix="dem_claim_ocr_") as tmp:
            text_variants = source_to_text_variants(source, Path(tmp))

    candidates = parse_candidates(text_variants)
    events, grouped = choose_event_values(candidates)

    missing = [
        spec["label"]
        for key, spec in LOAD_EVENTS.items()
        if spec.get("required", True) and key not in events
    ]
    if missing:
        print_report(events, grouped)
        raise RuntimeError(f"OCR could not find required LOAD events: {', '.join(missing)}")

    save_workbook(template, output, events, args)
    print_report(events, grouped)
    print(f"Saved workbook: {output}")


if __name__ == "__main__":
    main()
